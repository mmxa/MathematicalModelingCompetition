
"""
Description:data process program about Question1 and Question2,
including data pre-process and divided into segments
"""
import matplotlib.pyplot as plt
import xlrd
import re
from openpyxl import Workbook


class Node:
    def __init__(self):
        self.id = 0         # record segment id number
        self.vel = []       # record velocity for each step
        self.time_id = []   # record relative time for each step
        self.dt = []        # record time interval from GPS for each step
        self.ax = []        # calculate accelerate for each step
        self.state = []     # vehicle driving state: idle, smooth driving, accelerate, decelerate for each step
        self.ave_vel = 0    # calculate characteristic parameters for each segment
        self.ave_acc = 0
        self.ave_dec = 0
        self.ave_drive = 0
        self.drive_rate = 0
        self.idle_rate = 0
        self.acc_rate = 0
        self.dec_rate = 0
        self.distance = 0
        self.vel_standard_deviation = 0
        self.acc_standard_deviation = 0
        self.dec_standard_deviation = 0
        self.duration = 0
        self.max_vel = 0
        self.max_acc = 0
        self.max_dec = 0


class Solution:
    def __init__(self):
        self.rows = 0
        self.cols_time = []
        self.cols_vel = []
        self.cols_ax = []
        self.cols_ay = []
        self.cols_az = []
        self.cols_lat = []
        self.cols_lon = []
        self.wb = []
        self.start_time = []            # record relative time from primitive data，
        self.cols_rpm = []
        self.cols_tor = []
        self.cols_consumption = []
        self.cols_pedal = []
        self.cols_roil = []
        self.cols_rload = []
        self.cols_flow = []
        self.day = []
        self.dt = []                    # record time interval
        self.data_period = []           # first segment
        self.data_period2 = []          # segment record state for each step
        self.data_period3 = []          # segment dealing with idle state and irregular data
        self.data_period4 = []          # appropriate segment derived from former segments
        self.data_period5 = []          # sort id number and calculate parameters

    # import the data from excel
    def read_excel(self, filename):
        file = filename
        wr = xlrd.open_workbook(file)                      # open file
        sheet1 = wr.sheets()[0]
        self.rows = sheet1.nrows - 1                       # get number of lines from data?.xlsx
        self.cols_time = sheet1.col_values(0, 1)           # get contents of each column
        self.cols_vel = sheet1.col_values(1, 1)
        self.cols_ax = sheet1.col_values(2, 1)
        self.cols_ay = sheet1.col_values(3, 1)
        self.cols_az = sheet1.col_values(4, 1)
        self.cols_lat = sheet1.col_values(5, 1)
        self.cols_lon = sheet1.col_values(6,1)
        self.start_time.append(rematch(self.cols_time[0]))
        """
        # ignore some of data
        self.cols_rpm = sheet1.col_values(7,1)
        self.cols_tor = sheet1.col_values(8,1)
        self.cols_consumption = sheet1.col_values(9,1)
        self.cols_pedal = sheet1.col_values(10,1)
        self.cols_roil = sheet1.col_values(11,1)
        self.cols_rload = sheet1.col_values(12,1)
        self.cols_flow = sheet1.col_values(13,1)
        print(self.cols_time)
        """

    # data process program
    def process(self):
        for i in range(1, self.rows):
            self.start_time.append(rematch(self.cols_time[i]) - self.start_time[0])
        self.start_time[0] = 0
        for i in range(0, self.rows - 1):
            self.dt.append(self.start_time[i+1] - self.start_time[i])
        day = 1
        for i in range(0, self.rows - 1):
            if self.dt[i] > 20:
                day = day + 1
                self.day.append(day)            # record segment id, if interval from GPS greater than 20 second
            else:
                self.day.append(day)
        self.data_period = [Node() for i in range(day+1)]

        for i, day in enumerate(self.day):
            self.data_period[day].id = day
            self.data_period[day].vel.append(self.cols_vel[i])
            self.data_period[day].time_id.append(self.start_time[i])
            self.data_period[day].dt.append(self.dt[i])
        # print(self.data_period[1].vel)
        # print(self.data_period[1].dt)
        for i in range(self.day[len(self.day)-1] + 1):              # fill the segment from primitive data
            for j in range(len(self.data_period[i].vel)-1):
                self.data_period[i].ax.append((self.data_period[i].vel[j+1] -
                                               self.data_period[i].vel[j])/self.data_period[i].dt[j]/3.6)
        self.data_period2 = [Node() for i in range(day + 1)]
        for i in range(self.day[len(self.day)-1] + 1):
            for j in range(len(self.data_period[i].vel)-1):
                if self.data_period[i].dt[j] == 1:
                    self.data_period2[i].id = self.data_period[i].id
                    self.data_period2[i].vel.append(self.data_period[i].vel[j])
                    self.data_period2[i].time_id.append(self.data_period[i].time_id[j])
                    self.data_period2[i].dt.append(self.data_period[i].dt[j])
                    self.data_period2[i].ax.append(self.data_period[i].ax[j])
                    state = self.data_period[i].ax[j]
                    vel = self.data_period[i].vel[j]
                    if self.data_period[i].dt[j] > 20:                    # identify state of each step
                        self.data_period2[i].state.append("缺失")
                    elif abs(state) < 0.1 and vel >= 10:
                        self.data_period2[i].state.append("匀速")
                    elif state >= 0.1 and self.data_period[i].vel[j] > 10:
                        self.data_period2[i].state.append("加速")
                    elif state <= -0.1 and self.data_period[i].vel[j] > 10:
                        self.data_period2[i].state.append("减速")
                    else:
                        self.data_period2[i].state.append("怠速")
                elif self.data_period[i].dt[j] < 20:
                    for k in range(self.data_period[i].dt[j]):
                        self.data_period2[i].vel.append(self.data_period[i].vel[j] + k * self.data_period[i].ax[j])
                        self.data_period2[i].time_id.append(self.data_period[i].time_id[j] + k)
                        self.data_period2[i].dt.append(self.data_period[i].dt[j])
                        self.data_period2[i].ax.append(self.data_period[i].ax[j])
                        state = self.data_period[i].ax[j]
                        vel = self.data_period[i].vel[j]
                        if self.data_period[i].dt[j] > 20:
                            self.data_period2[i].state.append("缺失")
                        elif abs(state) < 0.1 and vel >= 10:
                            self.data_period2[i].state.append("匀速")
                        elif state >= 0.1 and self.data_period[i].vel[j] > 10:
                            self.data_period2[i].state.append("加速")
                        elif state <= -0.1 and self.data_period[i].vel[j] > 10:
                            self.data_period2[i].state.append("减速")
                        else:
                            self.data_period2[i].state.append("怠速")
                    """
                # ignore data with interval lager than 20 seconds
                else:
                    for k in range(self.data_period[i].dt[j]):
                        self.data_period2[i].vel.append(0)
                        self.data_period2[i].time_id.append(self.data_period[i].time_id[j] + k)
                        self.data_period2[i].dt.append(self.data_period[i].dt[j])
                        self.data_period2[i].ax.append(0)
                        state = self.data_period[i].ax[j]
                        vel = self.data_period[i].vel[j]
                        if self.data_period[i].dt[j] > 20:
                            self.data_period2[i].state.append("缺失")
                        elif abs(state) < 0.1 and vel >= 10:
                            self.data_period2[i].state.append("匀速")
                        elif state >= 0.1 and self.data_period[i].vel[j] > 3:
                            self.data_period2[i].state.append("加速")
                        elif state <= -0.1 and self.data_period[i].vel[j] > 3:
                            self.data_period2[i].state.append("减速")
                        else:
                            self.data_period2[i].state.append("怠速")
                    """
        # cut off idle state beyond 180 second
        self.data_period3 = [Node() for i in range(day + 1)]
        for i in range(self.day[len(self.day)-1] + 1):
            num = 0
            for j in range(len(self.data_period2[i].vel)-1):
                if self.data_period2[i].state[j] == '怠速' or self.data_period2[i].vel[j] < 3:
                    num = num + 1
                else:
                    num = 0
                if num > 180:
                    continue
                else:
                    self.data_period3[i].id = self.data_period2[i].id
                    self.data_period3[i].vel.append(self.data_period2[i].vel[j])
                    self.data_period3[i].time_id.append(self.data_period2[i].time_id[j])
                    self.data_period3[i].dt.append(self.data_period2[i].dt[j])
                    self.data_period3[i].ax.append(self.data_period2[i].ax[j])
                    self.data_period3[i].state.append(self.data_period2[i].state[j])
        print(len(self.data_period3))
        # generating fragments
        self.data_period4 = [Node() for i in range(10000)]
        k = 0
        for i in range(len(self.data_period3)):
            if len(self.data_period3[i].state) < 20:
                continue
            # new segment derived from next segment
            if self.data_period3[i].state[0] == '怠速':
                begin_flag = True
                end_flag = False
                record_flag = True
                k = k + 1
            else:
                begin_flag = False
                end_flag = False
                record_flag = False
            for j in range(len(self.data_period3[i].vel)-1):
                if self.data_period3[i].state[j] == '怠速' and begin_flag is True:
                    begin_flag = True
                    end_flag = False
                    record_flag = True
                if self.data_period3[i].state[j] == '加速':
                    begin_flag = False
                    end_flag = False
                # new segment derived from the same segment
                if self.data_period3[i].state[j] == '怠速' and begin_flag is False:
                    end_flag = True
                    k = k + 1
                    begin_flag = True
                    record_flag = True
                print(k)
                if record_flag is True:
                    self.data_period4[k].id = k
                    self.data_period4[k].vel.append(self.data_period3[i].vel[j])
                    self.data_period4[k].time_id.append(self.data_period3[i].time_id[j])
                    self.data_period4[k].dt.append(self.data_period3[i].dt[j])
                    self.data_period4[k].ax.append(self.data_period3[i].ax[j])
                    self.data_period4[k].state.append(self.data_period3[i].state[j])
            if end_flag is False:
                self.data_period4[k] = Node()           # clear the segment whose ending does not satisfied
                # self.data_period4[k].id = k
        # sort id number for each segment
        self.data_period5 = [Node() for i in range(2000)]
        k = 0
        for i in range(len(self.data_period4)):
            if len(self.data_period4[i].state) < 20:
                continue
            k = k + 1
            for j in range(len(self.data_period4[i].vel) - 1):
                self.data_period5[k].id = k
                self.data_period5[k].vel.append(self.data_period4[i].vel[j])
                self.data_period5[k].time_id.append(self.data_period4[i].time_id[j])
                self.data_period5[k].dt.append(self.data_period4[i].dt[j])
                self.data_period5[k].ax.append(self.data_period4[i].ax[j])
                self.data_period5[k].state.append(self.data_period4[i].state[j])
        # calculate characteristic parameters:
        for i in range(len(self.data_period5)):
            total_vel = 0
            total_acc = 0
            num_acc = 0.01
            num_dec = 0.01
            total_dec = 0
            total_drive = 0
            num_drive = 0.01
            num_idle = 0.01
            num = len(self.data_period5[i].vel)
            if num < 10:
                continue
            for j in range(len(self.data_period5[i].vel)):
                total_vel = total_vel + self.data_period5[i].vel[j]
                # num = num + 1
                if self.data_period5[i].state[j] == '加速':
                    num_acc = num_acc + 1
                    total_acc = total_acc + self.data_period5[i].ax[j]
                    total_drive = total_drive + self.data_period5[i].vel[j]
                elif self.data_period5[i].state[j] == '减速':
                    num_dec = num_dec + 1
                    total_dec = total_dec + self.data_period5[i].ax[j]
                    total_drive = total_drive + self.data_period5[i].vel[j]
                elif self.data_period5[i].state[j] == '匀速':
                    num_drive = num_drive + 1
                    total_drive = total_drive + self.data_period5[i].vel[j]
                else:
                    num_idle = num_idle + 1
            self.data_period5[i].ave_vel = total_vel / num
            self.data_period5[i].ave_acc = total_acc / num_acc
            self.data_period5[i].acc_rate = num_acc / num
            self.data_period5[i].ave_dec = total_dec / num_dec
            self.data_period5[i].dec_rate = num_dec / num
            self.data_period5[i].ave_dec = total_dec / num_dec
            self.data_period5[i].drive_rate = num_drive / num
            self.data_period5[i].ave_drive = total_drive / num_drive
            self.data_period5[i].idle_rate = num_idle / num
            self.data_period5[i].duration = num
            self.data_period5[i].distance = total_vel / 3.6
        # calculate characteristic parameters standard deviation
        for i in range(len(self.data_period5) - 1):
            total_vel_stand_deviation = 0
            total_acc_stand_deviation = 0
            total_dec_stand_deviation = 0
            num_acc = 0.01
            num_dec = 0.01
            max_vel = 0
            max_acc = 0
            max_dec = 0
            num = len(self.data_period5[i].vel)
            if num < 10:
                continue
            for j in range(len(self.data_period5[i].vel)):
                if self.data_period5[i].vel[j] > max_vel:
                    max_vel = self.data_period5[i].vel[j]
                total_vel_stand_deviation = total_vel_stand_deviation + \
                    (self.data_period5[i].vel[j]-self.data_period5[i].ave_vel)**2
                # self.data_period5[i]
                if self.data_period5[i].state[j] == '加速':
                    if self.data_period5[i].ax[j] > max_acc:
                        max_acc = self.data_period5[i].ax[j]
                    num_acc = num_acc + 1
                    total_acc_stand_deviation = total_acc_stand_deviation + \
                        (self.data_period5[i].ax[j] - self.data_period5[i].ave_acc) ** 2
                elif self.data_period5[i].state[j] == '减速':
                    if self.data_period5[i].ax[j] < max_dec:
                        max_dec = self.data_period5[i].ax[j]
                    num_dec = num_dec + 1
                    total_dec_stand_deviation = total_dec_stand_deviation + \
                        (self.data_period5[i].ax[j] - self.data_period5[i].ave_dec) ** 2
            self.data_period5[i].vel_standard_deviation = (total_vel_stand_deviation / num) ** 0.5
            self.data_period5[i].acc_standard_deviation = (total_acc_stand_deviation / num_acc) ** 0.5
            self.data_period5[i].dec_standard_deviation = (total_dec_stand_deviation / num_dec) ** 0.5
            self.data_period5[i].max_acc = max_acc
            self.data_period5[i].max_dec = max_dec
            self.data_period5[i].max_vel = max_vel

        """
        for i in range(1, self.day[len(self.day)-1] + 1):
            total_vel = 0
            total_acc = 0
            num_acc = 1
            num_dec = 1
            total_dec = 0
            total_drive = 0
            num_drive = 1
            num_idle = 1
            num = 1 
            for j in range(len(self.data_period3[i].vel)):
                total_vel = total_vel + self.data_period3[i].vel[j]
                num = num + 1
                if self.data_period3[i].state[j] == '加速':
                    num_acc = num_acc + 1
                    total_acc = total_acc + self.data_period3[i].ax[j]
                elif self.data_period3[i].state[j] == '减速':
                    num_dec = num_dec + 1
                    total_dec = total_dec + self.data_period3[i].ax[j]
                elif self.data_period3[i].state[j] == '匀速':
                    num_drive = num_drive + 1
                    total_drive = total_drive + self.data_period3[i].vel[j]
                else:
                    num_idle = num_idle + 1
            self.data_period3[i].ave_vel = total_vel / num
            self.data_period3[i].ave_acc = total_acc / num_acc
            self.data_period3[i].acc_rate = num_acc / num
            self.data_period3[i].ave_dec = total_dec / num_dec
            self.data_period3[i].dec_rate = num_dec / num
            self.data_period3[i].ave_dec = total_dec / num_dec
            self.data_period3[i].drive_rate = num_drive / num
            self.data_period3[i].ave_drive = total_drive / num_drive
            """
        # drawing the graph
    def figure(self):
        # fig = plt.figure()
        # fig, ax1 = plt.subplots()
        # temp = []
        # for i in range(1, self.rows):
        #    self.start_time.append(self.rematch(self.cols_time[i]) - self.start_time[0])
        # for i in range(1500,2001):
        #    temp.append((self.cols_vel[i+1]-self.cols_vel[i])/3.6-self.cols_ax[i]*9.8)
        # ax1 = plt.subplot(2, 2, 1)
        # ax1.plot(self.start_time[1:],linewidth=1)
        ax2 = plt.subplot(1, 1, 1)
        ax2.plot(self.data_period5[1].vel, 'r', linewidth=1)
        plt.show()
        """
        ax2.spines['top'].set_visible(False)
        ax2.spines['right'].set_visible(False)
        ax2.spines['bottom'].set_visible(False)
        ax2.spines['left'].set_visible(False)
        plt.xticks([])  # remove x axis
        plt.yticks([])  # remove y axis
        plt.axis('off')
        #
        ax1.plot(temp,c='r')
        ax1.set_xlabel("time")
        ax1.set_ylabel("velocity")
        ax2 = ax1.twinx()
        ax2.plot(self.start_time[1500:2000],self.cols_ax[1500:2000], 'b')
        plt.xlabel("time")
        ax2.set_ylabel("accerlation_x")"""
        """
        # figure a sketch of map
        ax3 = plt.subplot(2,2,3)
        for i in range(0, 20000):
            if self.day[i] % 6 == 0:
                ax3.plot([self.cols_lat[i], self.cols_lat[i+1]], [self.cols_lon[i],self.cols_lon[i+1]], 'b', linewidth=1)
            elif self.day[i] % 6 == 1:
                ax3.plot([self.cols_lat[i], self.cols_lat[i+1]], [self.cols_lon[i],self.cols_lon[i + 1]], 'r', linewidth=1)
            elif self.day[i] % 6 == 2:
                ax3.plot([self.cols_lat[i], self.cols_lat[i + 1]], [self.cols_lon[i], self.cols_lon[i + 1]], 'c', linewidth=1)
            elif self.day[i] % 6 == 3:
                ax3.plot([self.cols_lat[i], self.cols_lat[i+1]], [self.cols_lon[i],self.cols_lon[i + 1]], 'g', linewidth=1)
            elif self.day[i] % 6 == 4:
                ax3.plot([self.cols_lat[i], self.cols_lat[i+1]], [self.cols_lon[i],self.cols_lon[i + 1]], 'y', linewidth=1)
            elif self.day[i] % 6 == 5:
                ax3.plot([self.cols_lat[i], self.cols_lat[i+1]], [self.cols_lon[i],self.cols_lon[i + 1]], 'm', linewidth=1)
        """
    """
    def write_excle2(self,output):  # 将数据输出至EXCEL文件
        self.wb = xlwt.Workbook()
        sheet1 = self.wb.add_sheet('处理后', cell_overwrite_ok=True)
        # for i in range(0, 65000):
        # sheet1.write(i, 1, self.day[i])
        length = 0
        for i in range(10):
            for j in range(len(self.data_period[i + 1].ax)):
                sheet1.write(j + length, 2, self.data_period[i + 1].id)
                sheet1.write(j + length, 3, self.data_period[i + 1].time_id[j])
                sheet1.write(j + length, 4, self.data_period[i + 1].vel[j])
                sheet1.write(j + length, 5, self.data_period[i + 1].ax[j])
                length = length + len(self.data_period[i+1].time_id)
        self.wb.save(output)
    """

    # characteristic parameters Excel generating
    def write_excle(self, output):          # output data to excel
        outwb = Workbook()
        outws = outwb.create_sheet(index=0)
        length = 1
        outws.cell(1, 2).value = '序号'
        outws.cell(1, 3).value = '平均速度'
        outws.cell(1, 4).value = '平均加速度'
        outws.cell(1, 5).value = '加速度占比'
        outws.cell(1, 6).value = '平均减速度'
        outws.cell(1, 7).value = '减速度占比'
        outws.cell(1,8).value = '怠速占比'
        outws.cell(1, 9).value = '匀速占比'
        outws.cell(1, 10).value = '平均行驶车速'
        outws.cell(1, 11).value = '最大加速度'
        outws.cell(1, 12).value = '加速度标准差'
        outws.cell(1, 13).value = '最大减速度'
        outws.cell(1, 14).value = '减速度标准差'
        outws.cell(1, 15).value = '最大速度'
        outws.cell(1, 16).value = '速度标准差'
        outws.cell(1, 17).value = '行驶距离'
        outws.cell(1, 18).value = '行驶时长'
        for i in range(len(self.data_period5) - 1):
            if len(self.data_period5[i].vel) < 5:
                continue
            outws.cell(length + 1 + i, 2).value = self.data_period5[i].id
            # outws.cell(j + length + 1, 3).value = self.data_period4[i].time_id[j]
            # outws.cell(j + length + 1, 4).value = self.data_period4[i].vel[j]
            # outws.cell(j + length + 1, 5).value = self.data_period4[i].ax[j]
            # outws.cell(j + length + 1, 6).value = self.data_period4[i].dt[j]
            # outws.cell(j + length + 1, 7).value = self.data_period4[i].state[j]
            outws.cell(length + 1 + i, 3).value = self.data_period5[i].ave_vel
            outws.cell(length + 1 + i, 4).value = self.data_period5[i].ave_acc
            outws.cell(length + 1 + i, 5).value = self.data_period5[i].acc_rate
            outws.cell(length + 1 + i, 6).value = self.data_period5[i].ave_dec
            outws.cell( length + 1 + i, 7).value = self.data_period5[i].dec_rate
            outws.cell(length + 1 + i, 8).value = self.data_period5[i].idle_rate
            outws.cell(length + 1 + i, 9).value = self.data_period5[i].drive_rate
            outws.cell(length + 1 + i, 10).value = self.data_period5[i].ave_drive
            outws.cell(length + 1 + i, 11).value = self.data_period5[i].max_acc
            outws.cell(length + 1 + i, 12).value = self.data_period5[i].acc_standard_deviation
            outws.cell(length + 1 + i, 13).value = self.data_period5[i].max_dec
            outws.cell(length + 1 + i, 14).value = self.data_period5[i].dec_standard_deviation
            outws.cell(length + 1 + i, 15).value = self.data_period5[i].max_vel
            outws.cell(length + 1 + i, 16).value = self.data_period5[i].vel_standard_deviation
            outws.cell(length + 1 + i, 17).value = self.data_period5[i].distance
            outws.cell(length + 1 + i, 18).value = self.data_period5[i].duration

        outwb.save(output)

    # detial data Excel generating
    def write_excle_details(self, output):
        outwb = Workbook()
        outws = outwb.create_sheet(index=0)
        length = 1
        outws.cell(1, 2).value = '序号'
        outws.cell(1, 3).value = '相对时间'
        outws.cell(1, 4).value = '速度'
        outws.cell(1, 5).value = '加速度'
        outws.cell(1, 6).value = 'GPS时间间隔'
        outws.cell(1, 7).value = '运动状态'
        outws.cell(1, 8).value = '平均速度'
        outws.cell(1, 9).value = '平均加速度'
        outws.cell(1, 10).value = '加速度占比'
        outws.cell(1, 11).value = '平均减速度'
        outws.cell(1, 12).value = '减速度占比'
        outws.cell(1,13).value = '怠速占比'
        outws.cell(1, 14).value = '平均行驶车速'
        for i in range(len(self.data_period5) - 1):
            for j in range(len(self.data_period5[i].ax)):
                outws.cell(j + length + 1, 2).value = self.data_period5[i].id
                outws.cell(j + length + 1, 3).value = self.data_period5[i].time_id[j]
                outws.cell(j + length + 1, 4).value = self.data_period5[i].vel[j]
                outws.cell(j + length + 1, 5).value = self.data_period5[i].ax[j]
                outws.cell(j + length + 1, 6).value = self.data_period5[i].dt[j]
                outws.cell(j + length + 1, 7).value = self.data_period5[i].state[j]
                outws.cell(j + length + 1, 8).value = self.data_period5[i].ave_vel
                outws.cell(j + length + 1, 9).value = self.data_period5[i].ave_acc
                outws.cell(j + length + 1, 10).value = self.data_period5[i].acc_rate
                outws.cell(j + length + 1, 11).value = self.data_period5[i].ave_dec
                outws.cell(j + length + 1, 12).value = self.data_period5[i].dec_rate
                outws.cell(j + length + 1, 13).value = self.data_period5[i].idle_rate
                outws.cell(j + length + 1, 14).value = self.data_period5[i].ave_drive
                outws.cell(j + length + 1, 15).value = self.data_period5[i].duration
            length = length + len(self.data_period5[i].time_id)
        outwb.save(output)


# transform date to second
def rematch(temp):
    match_obj = re.search(r'(\d{4})/(\d\d)/(\d\d) (\d\d):(\d\d):(\d\d).(.*)', temp)
    day = int(match_obj.group(3))
    hour = int(match_obj.group(4))
    minute = int(match_obj.group(5))
    sec = int(match_obj.group(6))
    a = day * 3600 * 24 + hour * 3600 + 60 * minute + sec
    return a


def main():
    draw = False
    res = Solution()
    res.read_excel('data3.xlsx')
    res.process()
    res.write_excle('hehehe_tezheng333.xlsx')
    res.write_excle_details('hahaha_xiangqing333.xlsx')
    if draw is True:
        res.figure()
    print('Success! Congratulations !!!')


if __name__ == '__main__':
    main()
